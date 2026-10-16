"""
Build docs/finalists/HEDGE_Paper.docx from HEDGE_Paper.md plus src-images/.

APA 7 layout matched to the previous manuscript: Letter page, 1-inch margins,
Times New Roman 12 pt, double spacing, justified body with a 0.5-inch first-line
indent, centered bold Heading 1, flush-left bold Heading 2, flush-left bold
italic Heading 3, hanging-indent references.

Equations are native Word equations (OMML), not images. The table of contents
and every page cross-reference are real Word fields; word/settings.xml is given
<w:updateFields w:val="true"/> so Word resolves them when the file is opened.

Run:  python scratch/paper/build_docx.py
"""

import os
import re
import shutil
import sys
import zipfile

import docx
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SRC = os.path.join(REPO, "docs", "finalists", "HEDGE_Paper.md")
IMG = os.path.join(REPO, "docs", "finalists", "src-images")
OUT = os.path.join(REPO, "docs", "finalists", "HEDGE_Paper.docx")

FONT = "Times New Roman"
SIZE = Pt(12)
INDENT = Emu(457200)          # 0.5 inch
LETTER_W, LETTER_H = Emu(7772400), Emu(10058400)
MARGIN = Emu(914400)
BODY_W = Emu(7772400 - 2 * 914400)   # 6.5 inch

SCHOOL = ["RODOLFO V. FELICIANO MEMORIAL HIGH SCHOOL",
          "San Pedro II, Magalang, Pampanga"]
AUTHORS = ["Pasok, Charles Gideon T.", "Esguerra, Eireen Joy U.",
           "Valerio, Euan Deign D.", "Galang, Blessie M."]
TEACHER = ["Aleah G. Sabado, MAEd - Physical Science, LPT", "Research II Teacher"]

OMML_NS = ('xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
           'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"')


# ------------------------------------------------------------------ low level

def style_run(run, *, italic=False, bold=False, size=SIZE, mono=False):
    name = "Consolas" if mono else FONT
    run.font.name = name
    run.font.size = Pt(10) if mono else size
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), name)
    return run


def para(doc, text="", *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=None,
         space_after=0, bold=False, italic=False, size=SIZE, mono=False,
         hanging=False, left=None, spacing=2.0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing = spacing
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    if hanging:
        pf.left_indent = INDENT
        pf.first_line_indent = -INDENT
    else:
        if left is not None:
            pf.left_indent = left
        if first_line is not None:
            pf.first_line_indent = first_line
    if text:
        style_run(p.add_run(text), bold=bold, italic=italic, size=size, mono=mono)
    return p


def heading(doc, text, level):
    if level == 1:
        p = para(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    elif level == 2:
        p = para(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True)
    else:
        p = para(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, italic=True)
    # outline level so the TOC field can pick it up
    ppr = p._p.get_or_add_pPr()
    ol = OxmlElement("w:outlineLvl")
    ol.set(qn("w:val"), str(level - 1))
    ppr.append(ol)
    return p


def field(paragraph, instruction, placeholder="1"):
    """Insert a Word field code, e.g. 'TOC \\o "1-3" \\h \\z \\u' or 'PAGEREF x'."""
    r1 = paragraph.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r1._r.append(fc)
    style_run(r1)

    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = " " + instruction + " "
    r2._r.append(it)
    style_run(r2)

    r3 = paragraph.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "separate")
    r3._r.append(fc)
    style_run(r3)

    style_run(paragraph.add_run(placeholder))

    r5 = paragraph.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "end")
    r5._r.append(fc)
    style_run(r5)


_BM = [0]


def bookmark(paragraph, name):
    _BM[0] += 1
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(_BM[0]))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(_BM[0]))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def equation(doc, text):
    from docx.oxml import parse_xml
    esc = (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
    xml = (f'<m:oMath {OMML_NS}><m:r><w:rPr><w:rFonts w:ascii="{FONT}" '
           f'w:cs="{FONT}" w:eastAsia="{FONT}" w:hAnsi="{FONT}"/></w:rPr>'
           f'<m:t xml:space="preserve">{esc}</m:t></m:r></m:oMath>')
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    p._p.append(parse_xml(xml))
    return p


TABLE_W = 9360          # twips, 6.5 inch, matches the source manuscript
CELL_PT = Pt(10)


def _border(tag, single):
    e = OxmlElement("w:" + tag)
    e.set(qn("w:val"), "single" if single else "nil")
    e.set(qn("w:sz"), "4" if single else "0")
    e.set(qn("w:space"), "0")
    e.set(qn("w:color"), "000000")
    return e


def _cell_borders(cell, bottom=False):
    """APA tables carry no vertical rules and no interior horizontal rules."""
    tcpr = cell._tc.get_or_add_tcPr()
    old = tcpr.find(qn("w:tcBorders"))
    if old is not None:
        tcpr.remove(old)
    tb = OxmlElement("w:tcBorders")
    for tag in ("top", "left", "bottom", "right"):
        tb.append(_border(tag, single=(tag == "bottom" and bottom)))
    tcpr.append(tb)
    mar = OxmlElement("w:tcMar")
    for tag, w in (("top", 40), ("left", 80), ("bottom", 40), ("right", 80)):
        m = OxmlElement("w:" + tag)
        m.set(qn("w:w"), str(w))
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    tcpr.append(mar)


# Column widths in twips, carried over from the source manuscript so each
# rebuilt table keeps the proportions the previous one was laid out with.
# Every row sums to exactly TABLE_W.
APA_WIDTHS = {
    "Table 1":  [5200, 4160],
    "Table 2":  [3400, 1500, 4460],
    "Table 3":  [5200, 4160],
    "Table 4":  [3300, 600, 1100, 1500, 1400, 1460],
    "Table 5":  [2800, 1200, 1600, 1900, 1860],
    "Table 6":  [2800, 1200, 1600, 1900, 1860],
    "Table 7":  [5200, 4160],
    "Table 8":  [3400, 600, 1100, 1400, 1400, 1460],
    "Table 9":  [2090, 1272, 1727, 954, 954, 1181, 1182],
    "Table 10": [2090, 1272, 1727, 954, 954, 1181, 1182],
    "Table 11": [2090, 1272, 1727, 954, 954, 1181, 1182],
    "Table 12": [3000, 2200, 2200, 1960],
    "Table A1": [2100, 1200, 2200, 1400, 2460],
    "Table D1": [612, 656, 1476, 1476, 831, 503, 1093,
                 459, 459, 459, 459, 459, 418],
    "Table E1": [2340, 2340, 2340, 2340],
    "Table E2": [2600, 900, 900, 1500, 1500, 1960],
    "Table E3": [1808, 3085, 2234, 957, 1276],
}


def _column_widths(header, rows, total=TABLE_W):
    """Fixed layout needs explicit widths. Weight them by longest content."""
    n = len(header)
    weight = []
    for j in range(n):
        longest = max([len(header[j])] + [len(r[j]) for r in rows]) if rows else len(header[j])
        weight.append(max(longest, 4) ** 0.75)
    scale = total / sum(weight)
    w = [int(x * scale) for x in weight]
    floor = 420                     # about 0.29 inch, keeps flag columns legible
    w = [max(x, floor) for x in w]
    scale = total / sum(w)
    w = [int(x * scale) for x in w]
    w[-1] += total - sum(w)
    return w


def add_table(doc, header, rows, label=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = doc.styles["Normal Table"]
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.LEFT

    tblpr = t._tbl.tblPr
    for tag in ("w:tblW", "w:tblBorders", "w:tblLayout"):
        old = tblpr.find(qn(tag))
        if old is not None:
            tblpr.remove(old)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(TABLE_W))
    tw.set(qn("w:type"), "dxa")
    tblpr.append(tw)
    tb = OxmlElement("w:tblBorders")
    for tag in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tb.append(_border(tag, single=tag in ("top", "bottom")))
    tblpr.append(tb)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)

    widths = APA_WIDTHS.get(label) or _column_widths(header, rows)
    assert len(widths) == len(header), f"{label}: {len(widths)} widths for {len(header)} columns"
    assert sum(widths) == TABLE_W, f"{label}: widths sum to {sum(widths)}"

    # repeat the header row when a long table breaks across pages
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "1")
    trpr.append(th)

    def fill(cell, text, j, bold):
        _cell_borders(cell, bottom=bold)
        tcpr = cell._tc.get_or_add_tcPr()
        old = tcpr.find(qn("w:tcW"))
        if old is not None:
            tcpr.remove(old)
        cw = OxmlElement("w:tcW")
        cw.set(qn("w:w"), str(widths[j]))
        cw.set(qn("w:type"), "dxa")
        tcpr.insert(0, cw)
        cell.text = ""
        p = cell.paragraphs[0]
        pf = p.paragraph_format
        # first column reads as a stub, every other column is right-aligned,
        # which is what the source manuscript does
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
        pf.line_spacing = 1.0
        pf.space_before = Pt(1)
        pf.space_after = Pt(1)
        style_run(p.add_run(text), bold=bold, size=CELL_PT)

    for j, text in enumerate(header):
        fill(t.rows[0].cells[j], text, j, True)
    for i, row in enumerate(rows, start=1):
        for j, text in enumerate(row):
            fill(t.rows[i].cells[j], text, j, False)
    return t


def add_picture(doc, filename):
    path = os.path.join(IMG, filename)
    from PIL import Image
    with Image.open(path) as im:
        w, h = im.size
    height = Emu(int(BODY_W * h / w))
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(6)
    p.add_run().add_picture(path, width=BODY_W, height=height)


def keep_next(p):
    ppr = p._p.get_or_add_pPr()
    kn = OxmlElement("w:keepNext")
    kn.set(qn("w:val"), "1")
    ppr.append(kn)
    return p


def caption(doc, label, title):
    """APA 7 table/figure head: bold number line, then italic title line."""
    p = para(doc, label, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True,
             spacing=1.0, space_after=0)
    p.paragraph_format.space_before = Pt(10)
    keep_next(p)
    q = para(doc, title, align=WD_ALIGN_PARAGRAPH.LEFT, italic=True,
             spacing=1.0, space_after=4)
    keep_next(q)


def note_paragraph(doc, text):
    """A table or figure note: 10 pt italic throughout, as in the source paper.
    {PAGEREF name} becomes a live Word cross-reference field."""
    p = para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.0)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    for chunk in re.split(r"(\{PAGEREF \w+\})", text):
        if not chunk:
            continue
        m = re.fullmatch(r"\{PAGEREF (\w+)\}", chunk)
        if m:
            field(p, f"PAGEREF {m.group(1)} \\h")
        else:
            style_run(p.add_run(chunk), italic=True, size=Pt(10))
    return p


# ---------------------------------------------------------------- title page

def title_page(doc, title):
    for line in SCHOOL:
        para(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, spacing=1.5)
    logo = os.path.join(IMG, "school_logo.jpg")
    if os.path.exists(logo):
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(18)
        p.add_run().add_picture(logo, width=Emu(1097280))   # 1.2 inch
    para(doc, title, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, spacing=1.5,
         space_after=18)
    para(doc, "Research Paper", align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.5)
    para(doc, "By:", align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.5)
    for a in AUTHORS:
        para(doc, a, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.5)
    para(doc, "", spacing=1.5)
    for t in TEACHER:
        para(doc, t, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.5)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# --------------------------------------------------------------------- parse

def build():
    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = LETTER_W, LETTER_H
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, MARGIN)
    sec.header_distance = Emu(228600)      # 0.25 inch, as in the source paper
    sec.footer_distance = Emu(228600)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = SIZE

    # APA 7 page number, flush right, with a blank first page
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.line_spacing = 1.0
    hp.paragraph_format.space_after = Pt(0)
    field(hp, "PAGE")
    fp = sec.first_page_header.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.line_spacing = 1.0
    fp.paragraph_format.space_after = Pt(0)

    lines = open(SRC, encoding="utf-8").read().splitlines()
    i = 0
    title = ""
    n_fig = n_tbl = n_eq = 0

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        i += 1
        if not line:
            continue

        if line.startswith("@TITLE "):
            title = line[7:].strip()
            continue
        if line == "@TITLEPAGE":
            continue
        if line == "@PAGEBREAK":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            continue

        if line == "@TOC":
            heading(doc, "Table of Contents", 1)
            p = para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.5)
            field(p, 'TOC \\o "1-3" \\h \\z \\u',
                  "Right-click and choose Update Field to build the contents.")
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            continue

        if line.startswith("@H1 "):
            if title and line == "@H1 Abstract":
                title_page(doc, title)
            heading(doc, line[4:], 1)
            continue
        if line.startswith("@H2 "):
            heading(doc, line[4:], 2)
            continue
        if line.startswith("@H3 "):
            heading(doc, line[4:], 3)
            continue

        if line.startswith("@KEYWORDS "):
            p = para(doc, first_line=INDENT)
            style_run(p.add_run("Keywords: "), italic=True)
            style_run(p.add_run(line[len("@KEYWORDS Keywords: "):]))
            continue

        if line.startswith("@BOOKMARK "):
            name = line[10:].strip()
            nxt = lines[i].strip()
            assert nxt.startswith("@LIST "), nxt
            i += 1
            p = para(doc, nxt[6:], first_line=INDENT)
            bookmark(p, name)
            continue

        if line.startswith("@LIST "):
            para(doc, line[6:], left=INDENT, first_line=Emu(0), space_after=6)
            continue

        if line.startswith("@EQ "):
            n_eq += 1
            equation(doc, line[4:])
            continue

        if line.startswith("@MONO "):
            para(doc, line[6:], align=WD_ALIGN_PARAGRAPH.LEFT, mono=True,
                 spacing=1.0, space_after=6, left=INDENT)
            continue

        if line.startswith("@REF "):
            para(doc, line[5:], align=WD_ALIGN_PARAGRAPH.LEFT, hanging=True)
            continue

        if line.startswith("@FIG "):
            n_fig += 1
            _, cap, fname = [x.strip() for x in line[5:].split("|")]
            caption(doc, f"Figure {n_fig}", cap)
            add_picture(doc, fname)
            continue

        if line.startswith("@TABLE "):
            label, cap = [x.strip() for x in line[7:].split("|")]
            if label.isdigit():
                n_tbl += 1
                label = f"Table {n_tbl}"
                assert int(line[7:].split("|")[0].strip()) == n_tbl, line
            else:
                label = f"Table {label}"
            rows = []
            while lines[i].strip() != "@ENDTABLE":
                r = lines[i].strip()
                if r.startswith("|"):
                    rows.append([c.strip() for c in r.strip("|").split("|")])
                i += 1
            i += 1
            caption(doc, label, cap)
            add_table(doc, rows[0], rows[1:], label=label)
            continue

        if line.startswith("@NOTE "):
            note_paragraph(doc, line[6:])
            continue

        para(doc, line, first_line=INDENT)

    doc.save(OUT)
    print(f"wrote {OUT}")
    print(f"  figures {n_fig}  tables {n_tbl}  equations {n_eq}")
    return n_fig, n_tbl, n_eq


def set_update_fields(path):
    """Add <w:updateFields w:val='true'/> to word/settings.xml in place."""
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(
            tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/settings.xml":
                xml = data.decode("utf-8")
                if "updateFields" not in xml:
                    xml = xml.replace(
                        "<w:settings", "<w:settings", 1)
                    idx = xml.index(">", xml.index("<w:settings")) + 1
                    xml = xml[:idx] + '<w:updateFields w:val="true"/>' + xml[idx:]
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)
    print("  settings.xml: updateFields enabled")


if __name__ == "__main__":
    build()
    set_update_fields(OUT)
